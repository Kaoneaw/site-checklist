"""Reconcile embedded HDY data against the original workbook.

Developer-only check: python3 tests/verify_excel.py '/path/to/R3 Master Plan HDY TH.xlsx'
Requires openpyxl; opening the web app itself requires no dependencies.
"""
import collections
import datetime
import json
from pathlib import Path
import re
import sys
import openpyxl

html = (Path(__file__).resolve().parents[1] / 'index.html').read_text()
seed = json.loads(re.search(r'id="workbookSeed">(.*?)</script>', html, re.S)[1])
workbook = openpyxl.load_workbook(sys.argv[1], data_only=True)
assert len(seed['units']) == 13
assert len(seed['tasks']) == 185
assert sum(len(t['subtasks']) for t in seed['tasks']) == 632
assert sum(t['type'] == 'extra' for t in seed['tasks']) == 48
assert collections.Counter(t['status'] for t in seed['tasks']) == {0: 170, 1: 11, 2: 4}
assert sum(t['statusUnconfirmed'] for t in seed['tasks']) == 32
assert all(not t['completed'] and not t['owner'] and not t['requested'] for t in seed['tasks'])
assert all(t['amountKnown'] is False for t in seed['tasks'] if t['type'] == 'extra')
seen = set()
for task in seed['tasks']:
    sheet = workbook[task['sourceSheet']]
    row = task['sourceRow']
    assert task['name'] == str(sheet.cell(row, 2).value).strip()
    assert task['sourceStatus'] == str(sheet.cell(row, 3).value or '').strip()
    assert task['statusUnconfirmed'] == (not task['sourceStatus'])
    assert (sheet.title, row) not in seen
    seen.add((sheet.title, row))
    extra_header = next((r for r in range(8, row + 1)
                         if 'ตารางงานเพิ่มเติม' in str(sheet.cell(r, 1).value or '')), None)
    assert (task['type'] == 'extra') == (extra_header is not None)
    for sub in task['subtasks']:
        row = sub['sourceRow']
        assert sub['name'] == str(sheet.cell(row, 2).value).strip()
        assert sub['sourceStatus'] == str(sheet.cell(row, 3).value or '').strip()
        assert (sub['status'] is None) == (not sub['sourceStatus'])
        assert (sheet.title, row) not in seen
        seen.add((sheet.title, row))
        if sub['start']:
            assert sub['start'] == sheet.cell(row, 4).value.date().isoformat()
            weeks = int(re.fullmatch(r'(\d+)\s*สัปดาห์', sub['duration'])[1])
            end = datetime.date.fromisoformat(sub['start']) + datetime.timedelta(days=weeks * 7 - 1)
            assert sub['due'] == end.isoformat()
for sheet_name in {t['sourceSheet'] for t in seed['tasks']}:
    sheet = workbook[sheet_name]
    for row in range(8, sheet.max_row + 1):
        name = str(sheet.cell(row, 2).value or '').strip()
        if name and name != 'รายละเอียด / ขอบเขตงาน':
            assert (sheet_name, row) in seen, (sheet_name, row, name)
assert len(seen) == 817
print('PASS: 817 source rows accounted for exactly once; names, source status, missing values, Extra sections and explicit schedules match workbook')
